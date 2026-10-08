import json
from dataclasses import replace
from datetime import UTC, datetime
from pathlib import Path

import openpyxl
import pytest

from stock_candidate_report.export import export_report
from stock_candidate_report.models import Observation, ReportRow, RunResult, WatchItem


def test_snapshot_typed_values_and_formula_safety(tmp_path):
    row = ReportRow(
        WatchItem("AAA", "=1+1", "AI"),
        {
            "price": 100,
            "ytd": 0.25,
            "ntm_pe": 10,
            "eps_2026": 2,
            "revenue_2026": 100_000_000,
            "currency": "USD",
        },
        {},
    )
    result = RunResult("synthetic-run", [row], [], [], {"stocks": 1, "prices_available": 1})
    xlsx, manifest = export_report(result, 2026, tmp_path)
    data = json.loads(manifest.read_text(encoding="utf-8"))
    assert data["run_id"] == "synthetic-run"
    book = openpyxl.load_workbook(xlsx, data_only=False)
    assert book.sheetnames == ["增长和估值", "数据明细"]
    sheet = book["增长和估值"]
    assert sheet["C6"].data_type != "f"
    assert sheet["E6"].value == 100
    assert sheet["F6"].value == 0.25
    assert sheet["G6"].value == 10
    assert sheet["K6"].value is None
    assert sheet["M6"].value == 100
    assert sheet.freeze_panes is not None
    assert sheet.auto_filter.ref or len(sheet.tables) > 0
    assert "四个未来财季" in sheet["A3"].value
    again, _ = export_report(replace(result, run_id="synthetic-other"), 2026, tmp_path)
    assert xlsx != again


def test_refuse_same_run_overwrite_and_atomic_failure(tmp_path):
    result = RunResult("synthetic-run", [], [], [], {"stocks": 0, "prices_available": 0})
    export_report(result, 2026, tmp_path)
    with pytest.raises(FileExistsError):
        export_report(result, 2026, tmp_path)
    result = replace(result, run_id="will-fail")
    with pytest.raises(RuntimeError, match="artifact"):
        export_report(result, 2026, tmp_path, artifact_modules=Path("/missing-modules"))
    assert not (tmp_path / "候选股-will-fail.xlsx").exists()
    assert (tmp_path / "候选股-will-fail.json").exists()  # Diagnostic survives export failure.


def test_reject_unsafe_run_id(tmp_path):
    result = RunResult("../outside", [], [], [], {})
    with pytest.raises(ValueError, match="run_id"):
        export_report(result, 2026, tmp_path)


def test_quote_timestamp_is_visible_and_dates_are_typed(tmp_path):
    observation = Observation(
        "AAA",
        "price",
        100,
        "yahoo",
        datetime(2026, 10, 8, tzinfo=UTC),
        currency="USD",
        meta={"price_time": "2026-10-07T20:00:00+00:00"},
    )
    result = RunResult(
        "quote-time",
        [ReportRow(WatchItem("AAA"), {"price": 100}, {})],
        [observation],
        [],
        {"stocks": 1, "prices_available": 1},
    )
    output, _ = export_report(result, 2026, tmp_path)
    book = openpyxl.load_workbook(output, data_only=True)
    assert book["增长和估值"]["AC6"].value == datetime(2026, 10, 7, 20)
    assert book["数据明细"]["R6"].value == datetime(2026, 10, 7, 20)
